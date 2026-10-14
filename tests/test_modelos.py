"""Borradores 303 y 130 y libros registro sobre un trimestre de un profesional.

Escenario (ED simplificada, IVA general, 2026), importes calculados a mano:

1T
- Emitida a empresa: base 1.000, IVA 210, retención 150, total 1.060.
- Emitida a particular: base 500, IVA 105, total 605.
- Asesoría: base 100, IVA 21 (deducible).
- Ticket sin NIF: base 10, IVA 2,10 (no deducible → gasto 12,10).
- Gasolina: base 50, IVA 10,50 (IVA 50 % → 5,25; IRPF 0).
- SaaS de Irlanda (IE6388047V): base 20 sin IVA → autoliquidación 21 %.
- Ordenador: base 1.200, IVA 252 (inversión < 3.005,06 → casillas 28/29; IRPF 0).
- Gasto de 999 sin validar: fuera de todo.
2T
- Emitida a particular: base 1.000, IVA 210.
"""

from datetime import date
from decimal import Decimal as D

import pytest
from sqlalchemy import select

from aeat_hub.fiscal.libros import construir
from aeat_hub.fiscal.modelos import borrador_130, borrador_303, resumen_por_asiento
from aeat_hub.models import Actividad, Asiento, Titular

NIF = "12345678Z"


@pytest.fixture
def profesional(session):
    titular = Titular(nif=NIF, nombre="Ana Profesional")
    session.add(titular)
    session.flush()
    actividad = Actividad(
        codigo="AE-PRO-001", titular_id=titular.id, nombre="Consultoría", regimen="actividad_economica",
        regimen_irpf="estimacion_directa_simplificada", regimen_iva="general", epigrafe_iae="763",
        rendimiento_anterior=D("8000"),
    )
    session.add(actividad)
    session.flush()

    def alta(**kw):
        datos = {
            "actividad_id": actividad.id, "tipo": "gasto", "estado": "confirmado", "validado": True,
            "tipo_factura": "completa", "moneda": "EUR",
        }
        datos.update(kw)
        datos.setdefault("ejercicio", datos["fecha"].year)
        fila = Asiento(**datos)
        session.add(fila)
        session.flush()
        return fila

    ids = {}
    ids["emitida_empresa"] = alta(
        tipo="ingreso", cuenta_codigo="AE.ING.VENTAS", fecha=date(2026, 2, 1), nif_emisor=NIF, direccion="emitida",
        nif_receptor="B12345674", receptor="CLIENTE SL", numero_factura="2026-001", base=D("1000"),
        iva_tipo=D("21"), iva_cuota=D("210"), retencion_tipo=D("15"), retencion_cuota=D("150"), total=D("1060"),
    ).id
    ids["emitida_particular"] = alta(
        tipo="ingreso", cuenta_codigo="AE.ING.VENTAS", fecha=date(2026, 3, 1), nif_emisor=NIF,
        direccion="emitida", numero_factura="2026-002", base=D("500"), iva_tipo=D("21"), iva_cuota=D("105"),
        total=D("605"),
    ).id
    ids["asesoria"] = alta(
        cuenta_codigo="AE.GAS.PROFESIONALES", fecha=date(2026, 1, 15), nif_emisor="B84818442",
        nif_receptor=NIF, base=D("100"), iva_tipo=D("21"), iva_cuota=D("21"), total=D("121"),
    ).id
    ids["ticket"] = alta(
        cuenta_codigo="AE.GAS.SERVICIOS", fecha=date(2026, 1, 20), nif_emisor="A28812618",
        tipo_factura="simplificada", base=D("10"), iva_tipo=D("21"), iva_cuota=D("2.10"), total=D("12.10"),
    ).id
    ids["gasolina"] = alta(
        cuenta_codigo="AE.GAS.VEHICULO", fecha=date(2026, 2, 10), nif_emisor="A28141935", nif_receptor=NIF,
        base=D("50"), iva_tipo=D("21"), iva_cuota=D("10.50"), total=D("60.50"),
    ).id
    ids["saas"] = alta(
        cuenta_codigo="AE.GAS.SERVICIOS", fecha=date(2026, 3, 5), nif_emisor="IE6388047V", nif_receptor=NIF,
        base=D("20"), iva_cuota=None, total=D("20"),
    ).id
    ids["ordenador"] = alta(
        tipo="mejora", cuenta_codigo="AE.INV.BIENES", fecha=date(2026, 3, 20), nif_emisor="B12345674",
        nif_receptor=NIF, base=D("1200"), iva_tipo=D("21"), iva_cuota=D("252"), total=D("1452"),
    ).id
    ids["sin_validar"] = alta(
        cuenta_codigo="AE.GAS.SERVICIOS", fecha=date(2026, 3, 25), nif_emisor="B12345674", nif_receptor=NIF,
        base=D("999"), iva_tipo=D("21"), iva_cuota=D("209.79"), total=D("1208.79"), validado=False,
        estado="pendiente",
    ).id
    ids["emitida_2t"] = alta(
        tipo="ingreso", cuenta_codigo="AE.ING.VENTAS", fecha=date(2026, 5, 3), nif_emisor=NIF,
        direccion="emitida", numero_factura="2026-003", base=D("1000"), iva_tipo=D("21"), iva_cuota=D("210"),
        total=D("1210"),
    ).id
    session.commit()
    asientos = list(session.scalars(select(Asiento).where(Asiento.actividad_id == actividad.id)))
    return actividad, asientos, ids


def test_303_primer_trimestre(profesional):
    actividad, asientos, ids = profesional
    b = borrador_303(asientos, actividad, 2026, 1, titular_nif=NIF)
    assert b.aplica and b.verificado
    assert (b.importe("07"), b.importe("09")) == (D("1500.00"), D("315.00"))
    assert (b.importe("10"), b.importe("11")) == (D("20.00"), D("4.20"))
    assert b.importe("27") == D("319.20")
    assert (b.importe("28"), b.importe("29")) == (D("1325.00"), D("278.25"))
    assert (b.importe("36"), b.importe("37")) == (D("20.00"), D("4.20"))
    assert b.importe("45") == D("282.45")
    assert b.resultado == b.importe("46") == D("36.75")
    # trazabilidad: la cuota deducible sale de estos asientos y no del ticket ni del pendiente
    assert set(b.casillas["29"].asientos) == {ids["asesoria"], ids["gasolina"], ids["ordenador"]}
    assert ids["ticket"] not in resumen_por_asiento(b)
    assert b.n_por_revisar == 1 and b.total_por_revisar == D("1208.79")
    assert any("autoliquidación" in aviso for aviso in b.avisos)


def test_130_primer_y_segundo_trimestre(profesional):
    actividad, asientos, _ = profesional
    t1 = borrador_130(asientos, actividad, 2026, 1, titular_nif=NIF)
    assert t1.importe("01") == D("1500.00")
    # 100 + 12,10 + 0 + 20 + 0 = 132,10 ; difícil justificación 5 % de 1.367,90 = 68,40
    assert t1.importe("02") == D("200.50")
    assert t1.importe("04") == D("259.90")
    assert t1.importe("06") == D("150.00")
    assert t1.importe("07") == D("109.90")
    assert t1.importe("13") == D("100")  # rendimiento anterior 8.000 € ≤ 9.000
    assert t1.resultado == D("9.90")

    t2 = borrador_130(asientos, actividad, 2026, 2, titular_nif=NIF)
    assert t2.importe("01") == D("2500.00")
    assert t2.importe("02") == D("250.50")
    assert t2.importe("05") == D("109.90")
    assert t2.importe("07") == D("190.00")
    assert t2.resultado == D("90.00")


def test_capital_inmobiliario_no_presenta_303_ni_130(session):
    actividad = session.scalar(select(Actividad).where(Actividad.codigo == "CI-VA-001"))
    assert not borrador_303([], actividad, 2026, 1, titular_nif=None).aplica
    assert not borrador_130([], actividad, 2026, 1, titular_nif=None).aplica


def test_libros_registro(profesional):
    actividad, asientos, ids = profesional
    nombres = {"AE.ING.VENTAS": "Ingresos de explotación", "AE.GAS.SERVICIOS": "Otros servicios exteriores"}
    libros = construir(asientos, actividad, 2026, titular_nif=NIF, nombres_cuenta=nombres, trimestre=1)
    assert [a.numero for a in libros.ingresos] == ["2026-001", "2026-002"]
    assert libros.ingresos[0].nif == "B12345674"  # el destinatario en ventas
    assert libros.ingresos[0].importe_irpf == D("1000.00")
    assert libros.ingresos[0].retencion_cuota == D("150")
    gastos = {a.asiento_id: a for a in libros.gastos}
    assert ids["sin_validar"] not in gastos
    assert [a.numero_registro for a in libros.gastos] == list(range(1, len(libros.gastos) + 1))
    assert gastos[ids["ticket"]].importe_irpf == D("12.10")
    assert gastos[ids["gasolina"]].iva_deducible == D("5.25")
    assert [a.asiento_id for a in libros.bienes_inversion] == [ids["ordenador"]]
    assert libros.n_por_revisar == 1


def test_exportacion_excel_de_libros_y_borrador(session, layout, profesional):
    from openpyxl import load_workbook

    from aeat_hub.export_fiscal import borrador_de, libros_de

    _libros, ruta = libros_de(session, layout, "AE-PRO-001", 2026, 1)
    wb = load_workbook(ruta)
    assert wb.sheetnames == ["Ventas e ingresos", "Compras y gastos", "Bienes de inversión", "Notas"]
    assert wb["Ventas e ingresos"].max_row == 3  # cabecera + 2 facturas
    _b, ruta303 = borrador_de(session, layout, "303", "AE-PRO-001", 2026, 1)
    filas = {row[0]: row[2] for row in load_workbook(ruta303).active.iter_rows(min_row=2, values_only=True)}
    assert filas["46"] == 36.75


# --- Retenciones, informativas y resumen anual --------------------------------

from aeat_hub.fiscal.informativas import (  # noqa: E402
    borrador_111,
    borrador_115,
    borrador_347,
    borrador_349,
    resumen_iva_anual,
)


def _mas(session, actividad, **kw):
    datos = {"actividad_id": actividad.id, "tipo": "gasto", "estado": "confirmado", "validado": True,
             "tipo_factura": "completa", "moneda": "EUR", "nif_receptor": NIF}
    datos.update(kw)
    datos.setdefault("ejercicio", datos["fecha"].year)
    session.add(Asiento(**datos))
    session.commit()
    return list(session.scalars(select(Asiento).where(Asiento.actividad_id == actividad.id)))


def test_111_y_115_separan_profesionales_y_alquiler(session, profesional):
    actividad, _, _ = profesional
    _mas(session, actividad, cuenta_codigo="AE.GAS.PROFESIONALES", fecha=date(2026, 2, 2), nif_emisor="B84818442",
         base=D("1000"), iva_cuota=D("210"), retencion_tipo=D("15"), retencion_cuota=D("150"), total=D("1060"))
    asientos = _mas(session, actividad, cuenta_codigo="AE.GAS.ALQUILER", fecha=date(2026, 3, 1),
                    nif_emisor="12345678Z", base=D("800"), iva_cuota=D("168"), retencion_tipo=D("19"),
                    retencion_cuota=D("152"), total=D("816"))
    m111 = borrador_111(asientos, actividad, 2026, 1, titular_nif=NIF)
    assert (m111.importe("07"), m111.importe("08"), m111.importe("09")) == (D("1"), D("1000.00"), D("150.00"))
    m115 = borrador_115(asientos, actividad, 2026, 1, titular_nif=NIF)
    assert (m115.importe("01"), m115.importe("02"), m115.importe("03")) == (D("1"), D("800.00"), D("152.00"))


def test_capital_inmobiliario_no_retiene(session):
    actividad = session.scalar(select(Actividad).where(Actividad.codigo == "CI-VA-001"))
    assert not borrador_111([], actividad, 2026, 1).aplica


def test_347_umbral_y_exclusiones(session, profesional):
    actividad, _, _ = profesional
    _mas(session, actividad, cuenta_codigo="AE.GAS.COMPRAS", fecha=date(2026, 5, 2), nif_emisor="A28812618",
         emisor="PROVEEDOR SA", base=D("1652.89"), iva_cuota=D("347.11"), total=D("2000"))
    asientos = _mas(session, actividad, cuenta_codigo="AE.GAS.COMPRAS", fecha=date(2026, 8, 2),
                    nif_emisor="A28812618", base=D("826.45"), iva_cuota=D("173.55"), total=D("1000"))
    m347 = borrador_347(asientos, actividad, 2026)
    # A28812618: ticket 12,10 + 2.000 + 1.000 = 3.012,10 > 3.005,06
    assert m347.importe("A-A28812618") == D("3012.10")
    assert "1T 12.10" in m347.casillas["A-A28812618"].concepto
    # el cliente con retención va al 190 del cliente, no al 347; el ordenador no llega al umbral
    assert "B-B12345674" not in m347.casillas and "A-B12345674" not in m347.casillas


def test_349_saas_irlandes_como_servicio(profesional):
    actividad, asientos, ids = profesional
    m349 = borrador_349(asientos, actividad, 2026, 1)
    assert m349.importe("I-IE6388047V") == D("20.00")
    assert m349.casillas["I-IE6388047V"].asientos == [ids["saas"]]


def test_resumen_anual_suma_los_cuatro_trimestres(profesional):
    actividad, asientos, _ = profesional
    anual = resumen_iva_anual(asientos, actividad, 2026, titular_nif=NIF)
    assert anual.importe("09") == D("525.00")  # 315 (1T) + 210 (2T)
    assert anual.resultado == D("36.75") + D("210.00")
