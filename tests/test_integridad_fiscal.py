"""Regresiones P0 de integridad fiscal (auditoría 14 oct 2026).

Cada test reproduce un fallo que hacía desaparecer facturas reales del
libro o inventaba cifras. Los identificadores (A1, B3…) remiten al plan.
"""

from datetime import date
from decimal import Decimal
import json

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from aeat_hub.dedupe import DuplicateHit, find_duplicate
from aeat_hub.er import _candidatos
from aeat_hub.export import export_xlsx
from aeat_hub.extract.ids import normalize_numero
from aeat_hub.extract.parser import parse_invoice
from aeat_hub.extract.schema import InvoiceExtract
from aeat_hub.fiscal.cuadres import avisos_cuadre, cuadra_total, iva_tipo_conocido, recargo_tipo_conocido
from aeat_hub.fiscal.irpf import summarize_irpf
from aeat_hub.fiscal.money import last_amount, parse_amount
from aeat_hub.hub_http import _safe_upload_name
from aeat_hub.ingest import _apply_duplicate_and_state
from aeat_hub.models import Actividad, Asiento, Documento, Factura


def _actividad(session) -> Actividad:
    return session.scalar(select(Actividad).where(Actividad.codigo == "CI-VA-001"))


def _asiento(session, actividad, **kw) -> Asiento:
    data = {
        "actividad_id": actividad.id,
        "tipo": "gasto",
        "estado": "confirmado",
    }
    data.update(kw)
    if data.get("fecha") is not None and "ejercicio" not in kw:
        data["ejercicio"] = data["fecha"].year
    row = Asiento(**data)
    session.add(row)
    session.flush()
    return row


# --- A. Deduplicación -------------------------------------------------------


def test_a1_numeracion_reiniciada_por_ejercicio_no_agrupa(session):
    """Proveedor que reinicia la serie cada año: la '1' de 2026 no es la de 2025."""
    actividad = _actividad(session)
    session.add(
        Factura(
            actividad_id=actividad.id,
            nif_emisor="B12345674",
            numero_norm=normalize_numero("1"),
            numero_visible="1",
            fecha=date(2025, 1, 31),
            total=Decimal("700.00"),
        )
    )
    session.flush()
    extract = InvoiceExtract(
        nif_emisor="B12345674", numero="1", fecha=date(2026, 1, 31), total=Decimal("700.00")
    )
    assert _candidatos(session, actividad.id, extract) == []


def test_a1_mismo_ejercicio_si_agrupa(session):
    actividad = _actividad(session)
    session.add(
        Factura(
            actividad_id=actividad.id,
            nif_emisor="B12345674",
            numero_norm=normalize_numero("1"),
            fecha=date(2026, 1, 31),
            total=Decimal("700.00"),
        )
    )
    session.flush()
    extract = InvoiceExtract(
        nif_emisor="B12345674", numero="1", fecha=date(2026, 1, 31), total=Decimal("700.00")
    )
    assert len(_candidatos(session, actividad.id, extract)) == 1


def test_a2_mismo_numero_de_proveedores_distintos_no_es_duplicado(session):
    actividad = _actividad(session)
    _asiento(
        session,
        actividad,
        fecha=date(2026, 3, 1),
        nif_emisor="B12345674",
        emisor="FONTANERIA PEREZ SL",
        numero_factura="2026/0001",
        total=Decimal("121.00"),
    )
    extract = InvoiceExtract(
        nif_emisor="A28812618",
        emisor="ELECTRICIDAD GOMEZ SA",
        numero="2026-0001",
        fecha=date(2026, 3, 1),
        total=Decimal("121.00"),
    )
    assert find_duplicate(session, actividad_id=actividad.id, extract=extract, phash=None) is None


def test_a3_sospecha_sin_numero_no_se_excluye_sola():
    """Dos tickets iguales el mismo día (dos comidas): el segundo se revisa, no desaparece."""
    asiento = Asiento(tipo="gasto", estado="pendiente")
    hit = DuplicateHit(3, 1, 1, "mismo NIF+importe±3 días")
    _apply_duplicate_and_state(asiento, hit, Decimal("0.9"))
    assert asiento.estado == "pendiente"
    assert asiento.duplicado_de_id == 1
    assert asiento.duplicado_nivel == 3


def test_a4_rechazado_no_absorbe_reescaneo(session):
    actividad = _actividad(session)
    _asiento(
        session,
        actividad,
        fecha=date(2026, 4, 2),
        nif_emisor="B12345674",
        emisor="BAR PEPE",
        total=Decimal("12.50"),
        estado="rechazado",
        motivo_rechazo="OCR ilegible",
    )
    extract = InvoiceExtract(
        nif_emisor="B12345674", emisor="BAR PEPE", fecha=date(2026, 4, 2), total=Decimal("12.50")
    )
    assert find_duplicate(session, actividad_id=actividad.id, extract=extract, phash=None) is None


def test_a4_phash_no_cruza_expedientes_ni_rechazados(session):
    actividad = _actividad(session)
    otra = Actividad(
        titular_id=actividad.titular_id,
        codigo="OTRA-001",
        nombre="Otra",
        regimen=actividad.regimen,
    )
    session.add(otra)
    session.flush()
    doc = Documento(sha256="x" * 64, phash="f" * 64, nombre_original="a.jpg", ruta_almacenada="/a")
    session.add(doc)
    session.flush()
    _asiento(session, otra, documento_id=doc.id, total=Decimal("5.00"))
    extract = InvoiceExtract(total=Decimal("9.99"))
    assert find_duplicate(session, actividad_id=actividad.id, extract=extract, phash="f" * 64) is None

    doc2 = Documento(sha256="y" * 64, phash="e" * 64, nombre_original="b.jpg", ruta_almacenada="/b")
    session.add(doc2)
    session.flush()
    _asiento(session, actividad, documento_id=doc2.id, total=Decimal("5.00"), estado="rechazado")
    assert find_duplicate(session, actividad_id=actividad.id, extract=extract, phash="e" * 64) is None


# --- B. Importes, fechas e IVA ---------------------------------------------


@pytest.mark.parametrize(
    "raw, esperado",
    [
        ("-12,50", Decimal("-12.50")),
        ("12,50-", Decimal("-12.50")),
        ("(12,50)", Decimal("-12.50")),
        ("-1.234,56", Decimal("-1234.56")),
        ("1 234,56", Decimal("1234.56")),
        ("1 234,56", Decimal("1234.56")),
        ("1.234,5", Decimal("1234.50")),
        ("12,5", Decimal("12.50")),
        ("1.234,56 €", Decimal("1234.56")),
        ("48.40", Decimal("48.40")),
    ],
)
def test_b3_parse_amount_signos_y_miles(raw, esperado):
    assert parse_amount(raw) == esperado


def test_b3_punto_de_miles_sin_decimales_no_se_lee_como_decimal():
    # '1.234' es ambiguo (mil doscientos o uno con 234): nunca 1,23.
    assert parse_amount("1.234") != Decimal("1.23")


def test_b3_last_amount_conserva_el_signo_del_abono():
    assert last_amount("TOTAL ABONO -121,00 €") == Decimal("-121.00")
    assert last_amount("Total 1\u00a0234,56") == Decimal("1234.56")
    # Espacio normal: "2 345,50" es cantidad × precio, no dos mil.
    assert last_amount("2 345,50") == Decimal("345.50")


def _json(**data) -> str:
    return json.dumps(data)


def test_b1_no_inventa_iva_si_solo_hay_total():
    extract = parse_invoice(_json(emisor="MAPFRE", nif_emisor="A28141935", total=300.0))
    assert extract.total == Decimal("300.00")
    assert extract.iva_tipo is None
    assert extract.iva_cuota is None


def test_b2_no_retoca_la_base_para_cuadrar():
    extract = parse_invoice(
        _json(emisor="X", nif_emisor="B12345674", base=100.01, iva_cuota=21.0, total=121.0)
    )
    assert extract.base == Decimal("100.01")


def test_b4_fecha_no_iso_del_json_es_dia_mes():
    extract = parse_invoice(_json(emisor="X", nif_emisor="B12345674", fecha="03/02/2026", total=10))
    assert extract.fecha == date(2026, 2, 3)


def test_b4_fecha_iso_se_respeta():
    extract = parse_invoice(_json(emisor="X", nif_emisor="B12345674", fecha="2026-02-03", total=10))
    assert extract.fecha == date(2026, 2, 3)


def test_b5_importe_en_formato_espanol_como_cadena():
    extract = parse_invoice(
        _json(emisor="X", nif_emisor="B12345674", base="1.033,47", iva_cuota="217,03", total="1.250,50")
    )
    assert extract.total == Decimal("1250.50")
    assert extract.base == Decimal("1033.47")


def test_b6_cuadre_con_retencion_irpf():
    # Profesional: 100 base + 21 IVA − 15 IRPF = 106.
    assert cuadra_total(Decimal("100"), Decimal("21"), Decimal("106"), retencion=Decimal("15")) is True
    assert avisos_cuadre(
        base=Decimal("100"),
        iva_cuota=Decimal("21"),
        iva_tipo=Decimal("21"),
        total=Decimal("106"),
        retencion=Decimal("15"),
    ) == []


def test_b6_cuadre_con_recargo_de_equivalencia():
    # 100 base + 21 IVA + 5,20 RE = 126,20.
    assert cuadra_total(Decimal("100"), Decimal("21"), Decimal("126.20"), recargo=Decimal("5.20")) is True


@pytest.mark.parametrize("tipo", ["21", "10", "4", "0"])
def test_b6_tipos_generales(tipo):
    assert iva_tipo_conocido(Decimal(tipo), date(2026, 5, 1))


def test_b6_tipos_coyunturales_solo_en_su_ventana():
    assert iva_tipo_conocido(Decimal("5"), date(2023, 3, 1))
    assert iva_tipo_conocido(Decimal("2"), date(2024, 11, 1))
    assert not iva_tipo_conocido(Decimal("5"), date(2026, 3, 1))


@pytest.mark.parametrize("tipo", ["5.2", "1.4", "0.5", "1.75"])
def test_b6_tipos_de_recargo(tipo):
    assert recargo_tipo_conocido(Decimal(tipo))


def test_b6_tipo_inventado_sigue_siendo_raro():
    assert not iva_tipo_conocido(Decimal("17.3"))


# --- C. Agregados -----------------------------------------------------------


class _Row:
    def __init__(self, cuenta, total, estado="confirmado"):
        self.cuenta_codigo = cuenta
        self.total = total
        self.estado = estado


def test_c1_sin_clasificar_no_resta_como_gasto():
    resumen = summarize_irpf(
        [
            _Row("CI.ING.RENTA", Decimal("700.00")),
            _Row(None, Decimal("700.00")),  # recibo sin clasificar
        ]
    )
    assert resumen["gastos_deducibles"] == Decimal("0.00")
    assert resumen["rendimiento"] == Decimal("700.00")


def test_c1_rechazados_no_cuentan_en_irpf():
    resumen = summarize_irpf(
        [
            _Row("CI.ING.RENTA", Decimal("700.00")),
            _Row("CI.GAS.LUZ", Decimal("48.40"), estado="rechazado"),
        ]
    )
    assert resumen["gastos_deducibles"] == Decimal("0.00")


def test_c3_excel_gastos_no_incluye_duplicados(session, layout):
    actividad = _actividad(session)
    _asiento(
        session, actividad, fecha=date(2026, 3, 10), total=Decimal("48.40"), emisor="IBERDROLA", validado=True
    )
    _asiento(
        session,
        actividad,
        fecha=date(2026, 3, 10),
        total=Decimal("48.40"),
        emisor="IBERDROLA",
        estado="duplicado",
    )
    session.commit()
    path = export_xlsx(session, layout, actividad, 2026)
    wb = load_workbook(path)
    ws = wb["Gastos"]
    estados = [
        row[11]
        for row in ws.iter_rows(min_row=2, values_only=True)
        if row and row[8] is not None and isinstance(row[8], (int, float, Decimal))
    ]
    assert "duplicado" not in estados
    assert len(estados) == 1


# --- E. Seguridad -----------------------------------------------------------


@pytest.mark.parametrize(
    "raw, esperado",
    [
        ("factura.pdf", "factura.pdf"),
        ("../../etc/passwd.pdf", "passwd.pdf"),
        ("/tmp/x.jpg", "x.jpg"),
        ("..\\..\\win.png", "win.png"),
        ("ticket.exe", None),
        ("", None),
        ("..", None),
    ],
)
def test_e1_nombre_de_subida_saneado(raw, esperado):
    assert _safe_upload_name(raw) == esperado


# --- D. NIF -----------------------------------------------------------------


def test_d_nif_r_lleva_letra_de_control():
    from aeat_hub.fiscal.nif import is_valid_nif

    assert is_valid_nif("R2800123H")
    assert not is_valid_nif("R28001238")


@pytest.mark.parametrize("nif", ["K1234567L", "M7654321J"])
def test_d_nif_k_l_m_de_personas_sin_dni(nif):
    from aeat_hub.fiscal.nif import is_valid_nif

    assert is_valid_nif(nif)


@pytest.mark.parametrize(
    "texto, nif",
    [
        ("DNI: 12345678-Z", "12345678Z"),
        ("NIF 12.345.678-Z", "12345678Z"),
        ("CIF: B.12.345.674", "B12345674"),
        ("NIE X-1234567-L", "X1234567L"),
    ],
)
def test_d_nif_formatos_impresos(texto, nif):
    from aeat_hub.fiscal.nif import find_nifs

    assert find_nifs(texto) == [nif]


# --- E. Gemini --------------------------------------------------------------


def _gemini_ok(monkeypatch, tmp_path, body: dict):
    from PIL import Image

    from aeat_hub.ocr.gemini import GeminiOCRProvider
    from tests.samples import FACTURA_LUZ, write_pdf

    monkeypatch.setenv("GEMINI_API_KEY", "secreta")
    monkeypatch.setattr(
        "aeat_hub.ocr.gemini.iter_page_images",
        lambda path, dpi=180: iter([Image.new("RGB", (16, 16))]),
    )
    sent = {}

    class _Resp:
        def __enter__(self):
            return self

        def __exit__(self, *args):
            return False

        def read(self):
            return json.dumps(body).encode("utf-8")

    def _urlopen(req, timeout=75):
        sent["url"] = req.full_url
        sent["key"] = req.get_header("X-goog-api-key")
        return _Resp()

    monkeypatch.setattr("urllib.request.urlopen", _urlopen)
    pdf = write_pdf(tmp_path / "x.pdf", FACTURA_LUZ)
    return GeminiOCRProvider(), pdf, sent


def test_e2_api_key_en_cabecera_no_en_url(monkeypatch, tmp_path):
    body = {"candidates": [{"content": {"parts": [{"text": "{}"}]}, "finishReason": "STOP"}]}
    provider, pdf, sent = _gemini_ok(monkeypatch, tmp_path, body)
    provider.transcribe(pdf)
    assert "secreta" not in sent["url"]
    assert sent["key"] == "secreta"


def test_e2_respuesta_truncada_no_se_parsea(monkeypatch, tmp_path):
    from aeat_hub.ocr.base import ProviderUnavailable

    body = {
        "candidates": [
            {"content": {"parts": [{"text": '{"emisor": "X", "tot'}]}, "finishReason": "MAX_TOKENS"}
        ]
    }
    provider, pdf, _ = _gemini_ok(monkeypatch, tmp_path, body)
    with pytest.raises(ProviderUnavailable):
        provider.transcribe(pdf)


# --- A6. Deshacer una fusión sin pérdidas -----------------------------------


def test_a6_deshacer_fusion_restaura_numero_factura_y_evidencias(session):
    from aeat_hub.edits import desmarcar_duplicado, marcar_duplicado
    from aeat_hub.er import backfill_asientos, count_evidencias
    from aeat_hub.models import Cambio, Relacion

    actividad = _actividad(session)
    bueno = _asiento(
        session, actividad, fecha=date(2026, 5, 1), nif_emisor="B12345674",
        numero_factura="F-100", total=Decimal("50.00"), estado="confirmado",
    )
    gemelo = _asiento(
        session, actividad, fecha=date(2026, 5, 1), nif_emisor="B12345674",
        numero_factura="F-100-COPIA", total=Decimal("50.00"), estado="confirmado",
    )
    for row, sha in ((bueno, "a"), (gemelo, "b")):
        doc = Documento(sha256=sha * 64, nombre_original=f"{sha}.pdf", ruta_almacenada=f"/{sha}")
        session.add(doc)
        session.flush()
        row.documento_id = doc.id
    backfill_asientos(session)
    session.commit()
    factura_gemelo = gemelo.factura_id
    evid_antes = count_evidencias(session, factura_gemelo)
    assert evid_antes == 1

    marcar_duplicado(session, gemelo.id, bueno.id)
    session.commit()
    log_estado = session.scalar(
        select(Cambio).where(Cambio.asiento_id == gemelo.id, Cambio.campo == "estado")
    )
    assert log_estado.antes == "confirmado"  # no un "pendiente" inventado

    desmarcar_duplicado(session, gemelo.id)
    session.commit()
    session.expire_all()
    gemelo = session.get(Asiento, gemelo.id)
    assert gemelo.estado == "confirmado"
    assert gemelo.numero_factura == "F-100-COPIA"
    assert gemelo.factura_id == factura_gemelo
    assert gemelo.duplicado_de_id is None
    assert count_evidencias(session, factura_gemelo) == evid_antes
    assert not session.scalars(select(Relacion).where(Relacion.tipo == "misma_factura")).all()


# --- A7. Splink y facturas recurrentes ---------------------------------------


def test_a7_recurrentes_del_mismo_importe_no_son_duplicadas(session):
    """Alquiler mensual: mismo NIF, nombre e importe; números y meses distintos."""
    from aeat_hub.linkage_facturas import escanear_duplicados

    actividad = _actividad(session)
    for mes in (1, 2, 6):
        _asiento(
            session, actividad, fecha=date(2026, mes, 1), nif_emisor="B12345674",
            emisor="INMOBILIARIA SOL SL", numero_factura=f"2026-{mes:03d}",
            total=Decimal("700.00"),
        )
    session.commit()
    assert escanear_duplicados(session, actividad) == []


def test_a7_la_misma_factura_con_un_numero_mal_leido_si_se_sugiere(session):
    from aeat_hub.linkage_facturas import escanear_duplicados

    actividad = _actividad(session)
    for numero in ("2026-001", "2O26-0O1"):
        _asiento(
            session, actividad, fecha=date(2026, 3, 1), nif_emisor="B12345674",
            emisor="INMOBILIARIA SOL SL", numero_factura=numero, total=Decimal("700.00"),
        )
    session.commit()
    assert len(escanear_duplicados(session, actividad)) == 1


# --- B9. Edición en la ficha ---------------------------------------------------


def test_b9_editar_nif_invalido_se_rechaza(session):
    from aeat_hub.edits import patch_asiento

    row = _asiento(session, _actividad(session), total=Decimal("10.00"))
    with pytest.raises(ValueError, match="no es válido"):
        patch_asiento(session, row.id, {"nif_emisor": "B12345675"})


def test_b9_editar_nif_iva_ue_se_acepta(session):
    from aeat_hub.edits import patch_asiento

    row = _asiento(session, _actividad(session), total=Decimal("10.00"))
    patched = patch_asiento(session, row.id, {"nif_emisor": "IE6388047V"})
    assert patched.nif_emisor == "IE6388047V"


def test_b9_abono_escrito_con_coma_no_borra_el_importe(session):
    from aeat_hub.edits import patch_asiento

    row = _asiento(session, _actividad(session), total=Decimal("10.00"))
    patched = patch_asiento(session, row.id, {"total": "-12,50"})
    assert patched.total == Decimal("-12.50")


def test_b9_tipo_recalculado_se_ajusta_al_legal():
    from aeat_hub.fiscal.cuadres import tipo_iva_desde_cuota

    assert tipo_iva_desde_cuota(Decimal("0.95"), Decimal("0.20")) == Decimal("21.00")
    assert tipo_iva_desde_cuota(Decimal("100"), Decimal("17.30")) == Decimal("17.30")


# --- C4. Ranking de proveedores ----------------------------------------------


def test_c4_ranking_agrupa_por_nif_y_netea_abonos():
    from aeat_hub.dashboard import _rank_emisores

    filas = [
        {"estado": "confirmado", "validado": True, "tipo": "gasto", "total": Decimal("100"), "emisor": "IKEA Ibérica S.A.", "nif": "A28812618"},
        {"estado": "confirmado", "validado": True, "tipo": "gasto", "total": Decimal("50"), "emisor": "IKEA IBERICA SA", "nif": "A28812618"},
        {"estado": "confirmado", "validado": True, "tipo": "gasto", "total": Decimal("-30"), "emisor": "IKEA", "nif": "A28812618"},
        {"estado": "confirmado", "validado": True, "tipo": "gasto", "total": Decimal("40"), "emisor": "BAR", "nif": "—"},
    ]
    ranked, total = _rank_emisores(filas)
    assert total == Decimal("160")
    assert [r["total"] for r in ranked] == [Decimal("120"), Decimal("40")]


# --- Rectificativas -------------------------------------------------------------


def test_rectificativa_por_sustitucion_se_enlaza_al_ticket_sin_excluirlo(session):
    """Leroy Merlin: factura R… que sustituye al ticket, con prefijo repetido por el OCR."""
    from aeat_hub.ingest import vincular_rectificativa

    actividad = _actividad(session)
    ticket = _asiento(
        session, actividad, fecha=date(2026, 9, 5), nif_emisor="B84818442",
        numero_factura="064-0009-717875", total=Decimal("31.45"),
    )
    factura = _asiento(
        session, actividad, fecha=date(2026, 9, 6), nif_emisor="B84818442",
        numero_factura="064-0009-R194007", total=Decimal("31.45"),
        tipo_factura="rectificativa", factura_rectificada="064-064-0009-717875",
    )
    aviso = vincular_rectificativa(session, factura)
    assert factura.factura_rectificada_id == ticket.id
    assert "dos veces" in aviso
    assert ticket.estado == "confirmado"  # la decisión es humana


def test_original_que_llega_despues_tambien_se_enlaza(session):
    from aeat_hub.ingest import vincular_rectificativa

    actividad = _actividad(session)
    abono = _asiento(
        session, actividad, nif_emisor="B84818442", numero_factura="064-0009-R193996",
        total=Decimal("-10.99"), factura_rectificada="064-0009-R193994",
    )
    original = _asiento(
        session, actividad, nif_emisor="B84818442", numero_factura="064-0009-R193994",
        total=Decimal("10.99"),
    )
    assert vincular_rectificativa(session, original) is not None
    assert abono.factura_rectificada_id == original.id


def test_rectificativa_de_otro_emisor_no_se_enlaza(session):
    from aeat_hub.ingest import vincular_rectificativa

    actividad = _actividad(session)
    _asiento(session, actividad, nif_emisor="A28812618", numero_factura="F-000123", total=Decimal("5"))
    rect = _asiento(
        session, actividad, nif_emisor="B84818442", numero_factura="R-1",
        total=Decimal("-5"), factura_rectificada="F-000123",
    )
    assert vincular_rectificativa(session, rect) is None
    assert rect.factura_rectificada_id is None


# --- Seguridad: inyección de fórmulas en Excel -----------------------------------


def test_excel_no_ejecuta_formulas_del_ocr(session, layout):
    actividad = _actividad(session)
    malicioso = '=HYPERLINK("http://evil.example","pulsa")'
    _asiento(
        session, actividad, fecha=date(2026, 3, 10), total=Decimal("10.00"), emisor=malicioso,
        numero_factura="@SUM(1+1)", validado=True,
    )
    session.commit()
    ws = load_workbook(export_xlsx(session, layout, actividad, 2026))["Gastos"]
    celdas = [c for fila in ws.iter_rows(min_row=2) for c in fila if c.value in (malicioso, "@SUM(1+1)")]
    assert len(celdas) == 2
    assert all(c.data_type == "s" for c in celdas)
