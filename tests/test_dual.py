from datetime import date
from decimal import Decimal
from pathlib import Path

from aeat_hub.export_dual import write_dual_xlsx
from aeat_hub.extract.schema import InvoiceExtract
from aeat_hub.ocr.base import OCRResult
from aeat_hub.ocr.dual import compare_extracts, recommend, run_dual, vote_field
from aeat_hub.paths import DataLayout


def test_acuerdo_nif_e_importe():
    vote = vote_field("nif_emisor", "B-84818442", "B84818442", text_a="", text_b="")
    assert vote.estado == "acuerdo"
    money = vote_field("total", "13.86", "13.86", text_a="", text_b="")
    assert money.estado == "acuerdo"


def test_solo_un_motor_y_conflicto_no_promedia():
    solo = vote_field("numero", "064-0009-720394", "", text_a="", text_b="")
    assert solo.estado == "solo_vision"
    assert solo.consenso == "064-0009-720394"
    clash = vote_field("total", "13.86", "31.45", text_a="13,86", text_b="31,45")
    assert clash.estado == "conflicto"
    assert clash.consenso == ""


def test_nif_valido_gana_en_conflicto():
    vote = vote_field("nif_emisor", "B84818442", "B84818441", text_a="", text_b="")
    assert vote.estado == "conflicto"
    assert vote.consenso == "B84818442"


def test_recomendacion_revisar_si_total_choca():
    a = InvoiceExtract(nif_emisor="B84818442", numero="1", fecha=date(2026, 9, 18), total=Decimal("13.86"))
    b = InvoiceExtract(nif_emisor="B84818442", numero="1", fecha=date(2026, 9, 18), total=Decimal("31.45"))
    votes = compare_extracts(a, b, text_a="13,86", text_b="31,45")
    assert recommend(votes) == "revisar"


class _Fake:
    def __init__(self, name: str, text: str):
        self.name = name
        self._text = text

    def available(self):
        return True, "fake"

    def transcribe(self, path: Path) -> OCRResult:
        return OCRResult(text=self._text, engine=self.name, confidence=0.9)


LUZ = """\
IBERDROLA CLIENTES DEMO S.A.
NIF: B12345674
Factura nº: F2026-000123
Fecha: 10/03/2026
Cliente: 12345678Z
Base imponible: 40,00 €
IVA 21%: 8,40 €
Total factura: 48,40 €
"""


def test_run_dual_y_excel(layout: DataLayout, tmp_path: Path):
    pdf = tmp_path / "luz.pdf"
    pdf.write_text("placeholder")
    doc = run_dual(
        pdf,
        case_id="luz",
        vision=_Fake("apple-vision", LUZ),
        rapid=_Fake("rapidocr", LUZ),
    )
    assert doc.recomendacion == "consenso"
    assert doc.conflictos == 0
    dest = write_dual_xlsx(layout, [doc], dest=layout.exports / "dual_test.xlsx")
    assert dest.is_file()
    from openpyxl import load_workbook

    wb = load_workbook(dest)
    assert set(wb.sheetnames) >= {"Como_usarlo", "Resumen", "Campos", "Discrepancias", "Textos_OCR"}
    recs = [row[2].value for row in wb["Resumen"].iter_rows(min_row=2, max_col=3)]
    assert "consenso" in recs
