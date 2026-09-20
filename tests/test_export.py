from datetime import date
from decimal import Decimal

from openpyxl import load_workbook
from sqlalchemy import select

from aeat_hub.export import export_xlsx
from aeat_hub.models import Actividad, Asiento


def test_export_xlsx_crea_hojas(session, layout):
    actividad = session.scalar(select(Actividad).where(Actividad.codigo == "CI-VA-001"))
    session.add(
        Asiento(
            actividad_id=actividad.id,
            tipo="gasto",
            cuenta_codigo="CI.GAS.LUZ",
            fecha=date(2026, 3, 10),
            ejercicio=2026,
            emisor="IBERDROLA DEMO",
            nif_emisor="B12345674",
            numero_factura="F1",
            total=Decimal("48.40"),
            estado="confirmado",
        )
    )
    session.add(
        Asiento(
            actividad_id=actividad.id,
            tipo="ingreso",
            cuenta_codigo="CI.ING.RENTA",
            fecha=date(2026, 3, 1),
            ejercicio=2026,
            emisor="Inquilino",
            total=Decimal("700.00"),
            estado="confirmado",
        )
    )
    session.commit()
    dest = export_xlsx(session, layout, actividad, 2026)
    assert dest.is_file()
    wb = load_workbook(dest)
    assert set(wb.sheetnames) >= {
        "Gastos",
        "Ingresos",
        "Mejoras",
        "Duplicados",
        "Reclasificar",
        "Resumen_rubro",
        "Resumen_inmueble",
        "Casillas_IRPF",
    }
    gastos = list(wb["Gastos"].iter_rows(min_row=1, max_row=2, values_only=True))
    headers = list(gastos[0])
    assert "Emisor" in headers
    assert "Fecha compra" in headers
    emisor_idx = headers.index("Emisor")
    assert "Rubro" in headers
    assert "Rubro / cuenta" not in headers
    rubro_idx = headers.index("Rubro")
    assert gastos[1][emisor_idx] == "IBERDROLA DEMO"
    assert gastos[1][rubro_idx] == "Luz"
    ingresos_rows = list(wb["Ingresos"].iter_rows(min_row=1, values_only=True))
    rubro_idx = list(ingresos_rows[0]).index("Rubro")
    assert any(row[rubro_idx] == "Alquiler" for row in ingresos_rows[1:])
    for row in wb["Gastos"].iter_rows(min_row=2, values_only=True):
        assert not any(isinstance(cell, str) and cell.startswith("CI.") for cell in row)
    resumen = list(wb["Resumen_rubro"].iter_rows(min_row=1, values_only=True))
    assert resumen[0][0] == "rubro"
    assert any(row[0] == "Luz" for row in resumen[1:])
    irpf_rows = list(wb["Casillas_IRPF"].iter_rows(values_only=True))
    assert any(row and row[0] == "Rendimiento neto" for row in irpf_rows)
