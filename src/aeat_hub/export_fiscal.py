"""Exportación de libros registro y borradores de modelos a Excel."""

from __future__ import annotations

from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from aeat_hub.fiscal.libros import Anotacion, Libros, construir
from aeat_hub.fiscal.modelos import Borrador, borrador_130, borrador_303
from aeat_hub.models import Actividad, Asiento, Cuenta, Titular
from aeat_hub.paths import DataLayout
from aeat_hub.xlsx_seguro import guardar

FORMATO = "#,##0.00"
COLUMNAS_INGRESOS = (
    "Nº", "Fecha expedición", "Fecha operación", "Serie", "Número", "NIF destinatario", "Destinatario",
    "Concepto", "Base", "Tipo IVA", "Cuota IVA", "Tipo RE", "Cuota RE", "Total", "Tipo retención",
    "Retención", "Ingreso computable IRPF", "Rectificativa de",
)
COLUMNAS_GASTOS = (
    "Nº recepción", "Fecha expedición", "Fecha operación", "Serie", "Número", "NIF emisor", "Emisor",
    "Concepto", "Base", "Tipo IVA", "Cuota IVA", "Tipo RE", "Cuota RE", "Total", "Tipo retención",
    "Retención", "Cuota deducible", "Gasto deducible IRPF", "Rectificativa de", "Motivo",
)


def datos(session: Session, codigo: str):
    actividad = session.scalar(select(Actividad).where(Actividad.codigo == codigo))
    if actividad is None:
        raise RuntimeError(f"No existe la actividad {codigo}.")
    titular = session.get(Titular, actividad.titular_id)
    asientos = list(
        session.scalars(
            select(Asiento)
            .where(Asiento.actividad_id == actividad.id)
            .options(selectinload(Asiento.desglose_iva))
        )
    )
    nombres = {c.codigo: c.nombre for c in session.scalars(select(Cuenta))}
    return actividad, titular.nif if titular else None, asientos, nombres


def _filas(anotacion: Anotacion, ingreso: bool) -> list[list]:
    """Una fila por tipo impositivo (art. 2.d / 3.d de la Orden: base, tipo y cuota)."""
    desglose = anotacion.desglose or [None]
    filas = []
    for idx, iva in enumerate(desglose):
        primera = idx == 0
        comun = [
            anotacion.numero_registro, anotacion.fecha_expedicion, anotacion.fecha_operacion,
            anotacion.serie, anotacion.numero, anotacion.nif, anotacion.nombre, anotacion.concepto,
            iva.base if iva else None, iva.tipo if iva else None, iva.cuota if iva else None,
            iva.recargo_tipo if iva else None, iva.recargo_cuota if iva else None,
            anotacion.total if primera else None,
            anotacion.retencion_tipo if primera else None,
            anotacion.retencion_cuota if primera else None,
        ]
        if ingreso:
            filas.append(comun + [anotacion.importe_irpf if primera else None, anotacion.factura_rectificada])
        else:
            filas.append(
                comun
                + [
                    anotacion.iva_deducible if primera else None,
                    anotacion.importe_irpf if primera else None,
                    anotacion.factura_rectificada,
                    "; ".join(anotacion.reglas + anotacion.avisos) if primera else None,
                ]
            )
    return filas


def _hoja(ws, titulo: str, columnas, filas) -> None:
    ws.title = titulo
    ws.append(list(columnas))
    for celda in ws[1]:
        celda.font = Font(bold=True)
    for fila in filas:
        ws.append(fila)
    for row in ws.iter_rows(min_row=2):
        for celda in row:
            if isinstance(celda.value, Decimal):
                celda.number_format = FORMATO
            elif hasattr(celda.value, "isoformat"):
                celda.number_format = "DD/MM/YYYY"
    for idx, columna in enumerate(ws.columns, start=1):
        ancho = max((len(str(c.value or "")) for c in columna), default=8)
        ws.column_dimensions[get_column_letter(idx)].width = min(max(ancho + 2, 10), 50)
    ws.freeze_panes = "A2"


def exportar_libros(libros: Libros, destino: Path) -> Path:
    wb = Workbook()
    _hoja(wb.active, "Ventas e ingresos", COLUMNAS_INGRESOS,
          [f for a in libros.ingresos for f in _filas(a, True)])
    _hoja(wb.create_sheet(), "Compras y gastos", COLUMNAS_GASTOS,
          [f for a in libros.gastos for f in _filas(a, False)])
    _hoja(wb.create_sheet(), "Bienes de inversión", COLUMNAS_GASTOS,
          [f for a in libros.bienes_inversion for f in _filas(a, False)])
    notas = wb.create_sheet("Notas")
    notas.append(["Libros registro del IRPF (Orden HAC/773/2019), utilizables a efectos del IVA (art. 12)."])
    notas.append(["Solo incluye facturas validadas; duplicados y rechazadas fuera."])
    if libros.n_por_revisar:
        notas.append([f"Pendientes de revisar y NO incluidas: {libros.n_por_revisar} ({libros.total_por_revisar} €)."])
    for aviso in libros.avisos:
        notas.append([aviso])
    destino.parent.mkdir(parents=True, exist_ok=True)
    guardar(wb, destino)
    return destino


def exportar_borrador(borrador: Borrador, destino: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    _hoja(
        ws, f"Borrador {borrador.modelo}", ("Casilla", "Concepto", "Importe", "Asientos"),
        [
            [c.numero, c.concepto, c.importe, ", ".join(f"#{i}" for i in c.asientos)]
            for c in borrador.casillas.values()
        ],
    )
    notas = wb.create_sheet("Avisos")
    notas.append([f"Borrador del modelo {borrador.modelo} · {borrador.periodo} {borrador.ejercicio}. No es una presentación."])
    if not borrador.verificado:
        notas.append(["Parámetros del ejercicio NO verificados contra la AEAT."])
    for aviso in borrador.avisos:
        notas.append([aviso])
    destino.parent.mkdir(parents=True, exist_ok=True)
    guardar(wb, destino)
    return destino


def libros_de(session: Session, layout: DataLayout, codigo: str, year: int, trimestre: int | None) -> tuple[Libros, Path]:
    actividad, nif, asientos, nombres = datos(session, codigo)
    libros = construir(asientos, actividad, year, titular_nif=nif, nombres_cuenta=nombres, trimestre=trimestre)
    sufijo = f"_{trimestre}T" if trimestre else ""
    return libros, exportar_libros(libros, layout.exports / f"libros_{codigo}_{year}{sufijo}.xlsx")


def borrador_de(
    session: Session, layout: DataLayout, modelo: str, codigo: str, year: int, trimestre: int
) -> tuple[Borrador, Path]:
    actividad, nif, asientos, _ = datos(session, codigo)
    funcion = {"303": borrador_303, "130": borrador_130}.get(modelo)
    if funcion is None:
        raise RuntimeError("Modelos disponibles: 303 y 130.")
    borrador = funcion(asientos, actividad, year, trimestre, titular_nif=nif)
    return borrador, exportar_borrador(
        borrador, layout.exports / f"borrador_{modelo}_{codigo}_{year}_{trimestre}T.xlsx"
    )
