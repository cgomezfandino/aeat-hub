"""Excel de auditoría dual Vision vs RapidOCR. No es el libro fiscal."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from aeat_hub.ocr.dual import DualDocument
from aeat_hub.paths import DataLayout

HEADER = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")
FILL = {
    "acuerdo": PatternFill("solid", fgColor="C6EFCE"),
    "conflicto": PatternFill("solid", fgColor="FFC7CE"),
    "solo_vision": PatternFill("solid", fgColor="FFEB9C"),
    "solo_rapid": PatternFill("solid", fgColor="FFEB9C"),
    "vacio": PatternFill("solid", fgColor="D9D9D9"),
    "consenso": PatternFill("solid", fgColor="C6EFCE"),
    "revisar": PatternFill("solid", fgColor="FFC7CE"),
    "revisar_menor": PatternFill("solid", fgColor="FFEB9C"),
}


def write_dual_xlsx(layout: DataLayout, docs: list[DualDocument], dest: Path | None = None) -> Path:
    layout.exports.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    _uso(wb.active)
    _resumen(wb.create_sheet("Resumen"), docs)
    _campos(wb.create_sheet("Campos"), docs)
    _disc(wb.create_sheet("Discrepancias"), docs)
    _textos(wb.create_sheet("Textos_OCR"), docs)
    if dest is None:
        dest = layout.exports / "dual_vision_rapid.xlsx"
    wb.save(dest)
    return dest


def _uso(ws) -> None:
    ws.title = "Como_usarlo"
    lines = [
        ("Dual OCR Vision + RapidOCR", True),
        ("", False),
        ("Qué es", True),
        (
            "No sustituye al ingest ni a SQLite. Es una pasada de auditoría: "
            "los dos motores leen el mismo fichero, el parser saca los campos fiscales "
            "dos veces, y el Excel enseña acuerdos y conflictos.",
            False,
        ),
        ("", False),
        ("Flujo recomendado", True),
        ("1. Ingest normal (`aeat-hub ingest --ocr auto`) → SQLite + archivo/.", False),
        ("2. Dual sobre esas facturas (`aeat-hub dual`) → este Excel.", False),
        ("3. Revisa la hoja Discrepancias. Si hay conflicto en NIF/número/fecha/total, no te fíes del asiento.", False),
        ("4. Corrige en el libro (reclasificar / próximo ingest) y vuelve a exportar el libro fiscal.", False),
        ("", False),
        ("Reglas de consenso", True),
        ("acuerdo: ambos motores (tras normalizar NIF, importe ±0,01 €, número sin guiones) dicen lo mismo.", False),
        ("solo_vision / solo_rapid: un motor lo vio y el otro no. Se propone el que tiene valor.", False),
        ("conflicto: los dos dan valores distintos. No se promedian importes. NIF válido gana si el otro no lo es.", False),
        ("vacio: ninguno lo sacó (suele ser el parser, no el OCR).", False),
        ("", False),
        ("Recomendación por factura", True),
        ("consenso: campos críticos alineados. Más fiable que un solo motor.", False),
        ("revisar: conflicto o hueco en NIF emisor, número, fecha o total.", False),
        ("revisar_menor: conflicto en emisor/base/IVA, no en los cuatro críticos.", False),
        ("", False),
        ("Windows / Linux", True),
        ("Sin Apple Vision el dual no se puede completar: RapidOCR sigue siendo el motor portable del ingest.", False),
        ("Este Excel no se sube a git (exports/*.xlsx).", False),
    ]
    ws.column_dimensions["A"].width = 110
    for idx, (text, bold) in enumerate(lines, start=1):
        cell = ws.cell(idx, 1, text)
        cell.alignment = WRAP
        if bold:
            cell.font = HEADER
    ws.row_dimensions[4].height = 48
    ws.row_dimensions[8].height = 36


def _resumen(ws, docs: list[DualDocument]) -> None:
    headers = [
        "caso",
        "fichero",
        "recomendacion",
        "acuerdos",
        "conflictos",
        "huecos",
        "nif_consenso",
        "numero_consenso",
        "fecha_consenso",
        "total_consenso",
        "vision_ms",
        "rapid_ms",
        "vision_ok",
        "rapid_ok",
    ]
    _header_row(ws, headers)
    for row_i, doc in enumerate(docs, start=2):
        by = {item.campo: item for item in doc.votes}
        values = [
            doc.case_id,
            doc.path.name,
            doc.recomendacion,
            doc.acuerdos,
            doc.conflictos,
            doc.huecos,
            by.get("nif_emisor").consenso if "nif_emisor" in by else "",
            by.get("numero").consenso if "numero" in by else "",
            by.get("fecha").consenso if "fecha" in by else "",
            by.get("total").consenso if "total" in by else "",
            doc.vision.latency_ms,
            doc.rapid.latency_ms,
            "sí" if doc.vision.available else doc.vision.reason,
            "sí" if doc.rapid.available else doc.rapid.reason,
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row_i, col, value)
            if col == 3:
                cell.fill = FILL.get(str(value), PatternFill())
    _autosize(ws, len(headers))


def _campos(ws, docs: list[DualDocument]) -> None:
    headers = ["caso", "campo", "vision", "rapidocr", "consenso", "estado", "regla"]
    _header_row(ws, headers)
    row_i = 2
    for doc in docs:
        for vote in doc.votes:
            values = [doc.case_id, vote.campo, vote.vision, vote.rapid, vote.consenso, vote.estado, vote.regla]
            for col, value in enumerate(values, start=1):
                cell = ws.cell(row_i, col, value)
                if col == 6:
                    cell.fill = FILL.get(vote.estado, PatternFill())
            row_i += 1
    _autosize(ws, len(headers))


def _disc(ws, docs: list[DualDocument]) -> None:
    headers = ["caso", "campo", "vision", "rapidocr", "consenso", "estado", "regla"]
    _header_row(ws, headers)
    row_i = 2
    for doc in docs:
        for vote in doc.votes:
            if vote.estado in {"acuerdo"}:
                continue
            values = [doc.case_id, vote.campo, vote.vision, vote.rapid, vote.consenso, vote.estado, vote.regla]
            for col, value in enumerate(values, start=1):
                cell = ws.cell(row_i, col, value)
                if col == 6:
                    cell.fill = FILL.get(vote.estado, PatternFill())
            row_i += 1
    if row_i == 2:
        ws.cell(2, 1, "Sin discrepancias en este lote.")
    _autosize(ws, len(headers))


def _textos(ws, docs: list[DualDocument]) -> None:
    headers = ["caso", "motor", "ms", "caracteres", "texto (recorte 2500)"]
    _header_row(ws, headers)
    row_i = 2
    for doc in docs:
        for read in (doc.vision, doc.rapid):
            ws.cell(row_i, 1, doc.case_id)
            ws.cell(row_i, 2, read.engine)
            ws.cell(row_i, 3, read.latency_ms)
            ws.cell(row_i, 4, read.chars)
            cell = ws.cell(row_i, 5, (read.text or read.reason)[:2500])
            cell.alignment = WRAP
            ws.row_dimensions[row_i].height = 80
            row_i += 1
    ws.column_dimensions["E"].width = 80
    _autosize(ws, 4)


def _header_row(ws, headers: list[str]) -> None:
    for col, name in enumerate(headers, start=1):
        cell = ws.cell(1, col, name)
        cell.font = HEADER


def _autosize(ws, n_cols: int) -> None:
    for col in range(1, n_cols + 1):
        letter = get_column_letter(col)
        longest = 10
        for row in ws.iter_rows(min_col=col, max_col=col, max_row=min(ws.max_row, 40)):
            value = row[0].value
            if value is not None:
                longest = max(longest, min(len(str(value)), 48))
        ws.column_dimensions[letter].width = longest + 2
