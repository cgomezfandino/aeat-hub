"""Exportación de libros Excel por ejercicio."""

from __future__ import annotations

from collections import defaultdict
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from aeat_hub.fiscal.money import q2
from aeat_hub.fiscal.computo import computables, por_revisar
from aeat_hub.fiscal.irpf import summarize_irpf
from aeat_hub.fiscal.accounts import REGIMEN_CI
from aeat_hub.models import Actividad, Asiento, Cuenta, Inmueble
from aeat_hub.paths import DataLayout

SHEET_COLUMNS = [
    ("fecha", "Fecha compra"),
    ("emisor", "Emisor"),
    ("nif_emisor", "NIF emisor"),
    ("numero", "Nº factura"),
    ("descripcion", "Descripción"),
    ("base", "Base"),
    ("iva_tipo", "IVA %"),
    ("iva_cuota", "Cuota IVA"),
    ("total", "Total"),
    ("rubro", "Rubro"),
    ("tipo", "Tipo"),
    ("estado", "Estado"),
    ("inmueble", "Inmueble"),
    ("actividad", "Expediente"),
    ("id", "Id asiento"),
]


def export_xlsx(
    session: Session,
    layout: DataLayout,
    actividad: Actividad,
    year: int,
) -> Path:
    rows = session.scalars(
        select(Asiento)
        .where(Asiento.actividad_id == actividad.id, Asiento.ejercicio == year)
        .order_by(Asiento.fecha, Asiento.id)
    ).all()
    rows = [row for row in rows if row.estado != "rechazado"]
    filas_cuenta = session.scalars(select(Cuenta)).all()
    nombres = {item.codigo: item.nombre for item in filas_cuenta}
    extra_casillas = {item.codigo: item.casilla for item in filas_cuenta if item.casilla}
    inmuebles = {
        item.id: item.alias
        for item in session.scalars(select(Inmueble).where(Inmueble.actividad_id == actividad.id))
    }
    # Las hojas del libro solo llevan lo computable (validado, sin duplicados
    # ni rechazados): sumar una columna nunca cuenta dos veces ni datos sin mirar.
    libro = computables(rows)
    wb = Workbook()
    _write_sheet(wb.active, "Gastos", [r for r in libro if r.tipo == "gasto"], actividad, inmuebles, nombres)
    _write_sheet(
        wb.create_sheet("Ingresos"),
        "Ingresos",
        [r for r in libro if r.tipo == "ingreso"],
        actividad,
        inmuebles,
        nombres,
    )
    _write_sheet(
        wb.create_sheet("Mejoras"),
        "Mejoras",
        [r for r in libro if r.tipo == "mejora"],
        actividad,
        inmuebles,
        nombres,
    )
    _write_sheet(
        wb.create_sheet("Duplicados"),
        "Duplicados",
        [r for r in rows if r.estado == "duplicado"],
        actividad,
        inmuebles,
        nombres,
    )
    _write_sheet(
        wb.create_sheet("Por revisar"),
        "Por revisar",
        por_revisar(rows),
        actividad,
        inmuebles,
        nombres,
    )
    _write_trimestres(wb.create_sheet("Trimestres"), libro)
    _write_summary_rubro(wb.create_sheet("Resumen_rubro"), libro, nombres)
    _write_summary_inmueble(wb.create_sheet("Resumen_inmueble"), libro, inmuebles)
    if actividad.regimen == REGIMEN_CI:
        _write_casillas_irpf(wb.create_sheet("Casillas_IRPF"), libro, year, extra_casillas)

    layout.exports.mkdir(parents=True, exist_ok=True)
    dest = layout.exports / f"libro_{actividad.codigo}_{year}.xlsx"
    wb.save(dest)
    return dest


def _as_row(
    asiento: Asiento,
    actividad: Actividad,
    inmuebles: dict[int, str],
    nombres: dict[str, str],
) -> dict[str, object]:
    return {
        "id": asiento.id,
        "fecha": asiento.fecha.isoformat() if asiento.fecha else "",
        "emisor": asiento.emisor or "",
        "nif_emisor": asiento.nif_emisor or "",
        "numero": asiento.numero_factura or "",
        "descripcion": asiento.descripcion,
        "base": _num(asiento.base),
        "iva_tipo": _num(asiento.iva_tipo),
        "iva_cuota": _num(asiento.iva_cuota),
        "total": _num(asiento.total),
        "rubro": nombres.get(asiento.cuenta_codigo, "") if asiento.cuenta_codigo else "",
        "tipo": asiento.tipo,
        "estado": asiento.estado,
        "inmueble": inmuebles.get(asiento.inmueble_id or 0, ""),
        "actividad": actividad.codigo,
    }


def _write_sheet(
    ws,
    title: str,
    rows: list[Asiento],
    actividad: Actividad,
    inmuebles: dict[int, str],
    nombres: dict[str, str],
) -> None:
    ws.title = title
    ws.append([label for _key, label in SHEET_COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for asiento in rows:
        payload = _as_row(asiento, actividad, inmuebles, nombres)
        ws.append([payload[key] for key, _label in SHEET_COLUMNS])
    _autosize(ws)


def _write_trimestres(ws, rows: list[Asiento]) -> None:
    """Cortes T1–T4 sobre el libro anual: base del borrador 303 (sin presentación)."""
    ws.title = "Trimestres"
    ws.append(["Trimestre", "Gastos", "Ingresos", "Mejoras", "Neto", "IVA soportado"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    trimestres = {
        q: {"gasto": Decimal("0"), "ingreso": Decimal("0"), "mejora": Decimal("0"), "iva": Decimal("0")}
        for q in (1, 2, 3, 4)
    }
    totales = {"gasto": Decimal("0"), "ingreso": Decimal("0"), "mejora": Decimal("0"), "iva": Decimal("0")}
    for asiento in rows:
        if asiento.estado == "duplicado" or asiento.fecha is None:
            continue
        slot = trimestres[min((asiento.fecha.month - 1) // 3 + 1, 4)]
        if asiento.tipo in ("gasto", "ingreso", "mejora"):
            slot[asiento.tipo] += asiento.total or Decimal("0")
            totales[asiento.tipo] += asiento.total or Decimal("0")
        if asiento.tipo in ("gasto", "mejora"):
            slot["iva"] += asiento.iva_cuota or Decimal("0")
            totales["iva"] += asiento.iva_cuota or Decimal("0")
    for q in (1, 2, 3, 4):
        slot = trimestres[q]
        ws.append(
            [
                f"T{q}",
                _num(slot["gasto"]),
                _num(slot["ingreso"]),
                _num(slot["mejora"]),
                _num(slot["ingreso"] - slot["gasto"]),
                _num(slot["iva"]),
            ]
        )
    ws.append(
        [
            "Año",
            _num(totales["gasto"]),
            _num(totales["ingreso"]),
            _num(totales["mejora"]),
            _num(totales["ingreso"] - totales["gasto"]),
            _num(totales["iva"]),
        ]
    )
    _autosize(ws)


def _write_summary_rubro(ws, rows: list[Asiento], nombres: dict[str, str]) -> None:
    ws.title = "Resumen_rubro"
    ws.append(["rubro", "tipo", "n_asientos", "base", "iva", "total"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    buckets: dict[tuple[str, str], dict[str, Decimal | int]] = defaultdict(
        lambda: {"n": 0, "base": Decimal("0"), "iva": Decimal("0"), "total": Decimal("0")}
    )
    for asiento in rows:
        if asiento.estado == "duplicado":
            continue
        rubro = nombres.get(asiento.cuenta_codigo, "(sin rubro)") if asiento.cuenta_codigo else "(sin rubro)"
        key = (rubro, asiento.tipo)
        buckets[key]["n"] += 1
        buckets[key]["base"] += asiento.base or Decimal("0")
        buckets[key]["iva"] += asiento.iva_cuota or Decimal("0")
        buckets[key]["total"] += asiento.total or Decimal("0")
    for (rubro, tipo), agg in sorted(buckets.items()):
        ws.append([rubro, tipo, agg["n"], _num(agg["base"]), _num(agg["iva"]), _num(agg["total"])])
    _autosize(ws)


def _write_summary_inmueble(ws, rows: list[Asiento], inmuebles: dict[int, str]) -> None:
    ws.title = "Resumen_inmueble"
    ws.append(["inmueble", "tipo", "total"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    buckets: dict[tuple[str, str], Decimal] = defaultdict(lambda: Decimal("0"))
    for asiento in rows:
        if asiento.estado == "duplicado":
            continue
        alias = inmuebles.get(asiento.inmueble_id or 0, "(sin inmueble)")
        buckets[(alias, asiento.tipo)] += asiento.total or Decimal("0")
    for (alias, tipo), total in sorted(buckets.items()):
        ws.append([alias, tipo, _num(total)])
    _autosize(ws)


def _write_casillas_irpf(
    ws,
    rows: list[Asiento],
    year: int,
    extra_casillas: dict[str, str],
) -> None:
    ws.title = "Casillas_IRPF"
    summary = summarize_irpf(rows, extra_casillas)
    ws.append(["Borrador IRPF · rendimientos de capital inmobiliario", year])
    ws.append(
        [
            "No presenta el modelo 100. Copia estos totales al anexo del inmueble en la Renta.",
        ]
    )
    ws.append([])
    ws.append(["Concepto", "Asientos", "Importe", "Resta del año", "Nota"])
    for cell in ws[4]:
        cell.font = Font(bold=True)
    for item in summary["filas"]:
        if item["n"] == 0 and item["clave"] not in {"ingresos", "mejoras"}:
            continue
        ws.append(
            [
                item["etiqueta"],
                item["n"],
                _num(item["total"]),
                "Sí" if item["resta_del_ano"] else "No",
                item["notas"],
            ]
        )
    ws.append([])
    ws.append(["Ingresos íntegros", _num(summary["ingresos"])])
    ws.append(["Gastos deducibles (incl. amortización)", _num(summary["gastos_deducibles"])])
    ws.append(["Rendimiento neto", _num(summary["rendimiento"])])
    ws.append(["Mejoras (no restan del ejercicio)", _num(summary["mejoras"])])
    _autosize(ws)


def _num(value) -> Decimal | None:
    """Importe exacto para Excel (Decimal, no float binario)."""
    if value is None:
        return None
    return q2(value)


FORMATO_IMPORTE = "#,##0.00"


def _autosize(ws) -> None:
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, Decimal):
                cell.number_format = FORMATO_IMPORTE
    for idx, column in enumerate(ws.columns, start=1):
        width = max((len(str(cell.value or "")) for cell in column), default=8)
        ws.column_dimensions[get_column_letter(idx)].width = min(max(width + 2, 10), 40)
